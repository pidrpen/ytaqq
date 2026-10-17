"""Integration test against an already running LOCAL llama-server.
Build agent_probe.c first. All file operations stay in a new test directory.
python native/test_agent_model.py --probe ./agent_probe.exe --scratch ../model-tests
"""
import argparse
import json
from pathlib import Path
import subprocess
import time
import urllib.request
import zipfile
import xml.etree.ElementTree as ET


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--probe', required=True, type=Path)
    parser.add_argument('--scratch', required=True, type=Path)
    parser.add_argument('--port', type=int, default=39285)
    parser.add_argument('--case', action='append', dest='only')
    args = parser.parse_args()
    probe = str(args.probe.resolve())
    system = json.loads(subprocess.check_output([probe, '--system'], encoding='utf-8'))
    schema = json.loads(subprocess.check_output([probe, '--schema'], encoding='utf-8'))
    args.scratch.mkdir(parents=True, exist_ok=False)
    cases = [
        ('excel-formula', 'Создай на рабочем столе книгу Расчёт.xlsx, запиши числа 12 и 3 и формулой SUM посчитай их сумму.', {}, None),
        ('excel-path', '{root}\\Работа создай ексель', {}, None),
        ('excel-zero', 'с нуля создай ексель', {}, None),
        ('excel-clarify', '{root}\\Работа создай ексель\nУточнение пользователя: с нуля создай ексель', {}, None),
        ('excel-word', 'Создай на рабочем столе в папке Офис пустой эксель и ворд.', {}, None),
        ('word-there', 'создай там ворд', {}, None),
        ('word-stl', 'создай папку на рабочем столе и создай в нем ворд файл, и в ту же папку stl файлы с рабочего стола кинь',
         {'Деталь.stl': 'a', 'переходник.stl': 'b', 'подсттава.stl': 'c', 'keep.stl.txt': 'keep'}, None),
        ('inspect-34d', 'какие файлы лежат в папке 34д на рабочем столе?',
         {**{f'34д/документ-{i:02d}.txt': f'text {i}' for i in range(70)}, 'outside.stl': 'outside'}, None),
        ('stl', 'нужно создать на рабочем столе папку Модели и засунуть туда все файлы .stl',
         {'деталь.stl': 'a', 'КОРПУС.STL': 'b', 'деталь.stl.txt': 'keep', 'stl-note.txt': 'keep'},
         {'Модели/деталь.stl', 'Модели/КОРПУС.STL', 'деталь.stl.txt', 'stl-note.txt'}),
        ('multi', 'На рабочем столе перемести все .stl в папку Модели, а все .pdf скопируй в папку Архив на рабочем столе.',
         {'part.stl': 'a', 'plan.pdf': 'b', 'leave.txt': 'c'},
         {'Модели/part.stl', 'Архив/plan.pdf', 'plan.pdf', 'leave.txt'}),
        ('new-extension', 'С рабочего стола собери все файлы .3mf в новую папку Печать на рабочем столе, перенеси их туда.',
         {'part.3mf': 'a', 'PART2.3MF': 'b', 'part.3mf.txt': 'c'},
         {'Печать/part.3mf', 'Печать/PART2.3MF', 'part.3mf.txt'}),
        ('rename', 'Переименуй на рабочем столе файл old.stl в корпус.stl.',
         {'old.stl': 'a', 'keep.txt': 'b'}, {'корпус.stl', 'keep.txt'}),
        ('list', 'Покажи, какие файлы сейчас на рабочем столе?',
         {'реальный-файл.txt': 'a'}, {'реальный-файл.txt'}),
        ('delete', 'Удали все файлы с рабочего стола.',
         {'keep.stl': 'a'}, {'keep.stl'}),
    ]
    report = []
    for label, goal, files, expected in cases:
        if args.only and label not in args.only:
            continue
        root = (args.scratch / label).resolve()
        root.mkdir()
        goal = goal.replace('{root}', str(root))
        for name, content in files.items():
            (root / name).parent.mkdir(parents=True, exist_ok=True)
            (root / name).write_text(content, encoding='utf-8')
        process = subprocess.Popen([probe, str(root)], stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                   text=True, encoding='utf-8')

        def tool(text):
            process.stdin.write(text + '\n')
            process.stdin.flush()
            return json.loads(process.stdout.readline())

        tool(json.dumps({'goal': goal}, ensure_ascii=False))
        assert subprocess.check_output([probe, '--route', goal], text=True).strip() == '1', f'routing failed: {label}'
        plan, notes, log = [], '', []
        state = None
        errors = 0
        started = time.time()
        try:
            for step in range(1, 17):
                prompt = (f'Задача пользователя: {goal}\nПапки: @desktop={root}; '
                          f'@downloads={root}\\Downloads; @documents={root}\\Documents; @last={root}.\n'
                          f'Шаг {step} из 16. Действий в плане: {len(plan)}.\n')
                kinds = ['mkdir', 'copy', 'move', 'rename', 'open', 'write', 'write', 'write', 'excel_write', 'inventory']
                for action in plan[:12]:
                    prompt += f'В плане: {kinds[action["kind"]]} {action["src"]} -> {action["dst"]}\n'
                prompt += notes[-12000:]
                body = json.dumps({'model': 'local', **schema, 'messages': [
                    {'role': 'system', 'content': system}, {'role': 'user', 'content': prompt}],
                    'temperature': 0.2, 'max_tokens': 512, 'stream': False}).encode('utf-8')
                req = urllib.request.Request(f'http://127.0.0.1:{args.port}/v1/chat/completions', data=body,
                                             headers={'Content-Type': 'application/json'})
                with urllib.request.urlopen(req, timeout=180) as response:
                    answer = json.load(response)['choices'][0]['message']['content']
                result = tool(answer.replace('\n', ' '))
                state, plan = result['state'], result['plan']
                log.append({'step': step, 'model': answer, 'result': result})
                print(f'{label}: step={step} state={state} count={len(plan)} {answer[:180]}', flush=True)
                notes += f'\nВызов: {answer}\nРезультат: {result["observation"]}\n'
                if state in (1, 2):
                    break
                if state < 0:
                    errors += 1
                    if errors >= 3:
                        break
            assert state in (1, 2), 'agent did not finish'
            before = {p.relative_to(root).as_posix() for p in root.rglob('*') if p.is_file()}
            assert before == set(files), 'preview changed files'
            if label == 'delete':
                assert state == 2 and not plan, 'delete must be refused'
            elif label == 'inspect-34d':
                assert not plan and all(f'документ-{i:02d}.txt' in notes for i in range(70)), 'incomplete actual folder listing'
                assert not any('outside.stl' in entry['result']['observation'] for entry in log), 'wrong folder listed'
            elif label == 'list':
                assert not plan and 'реальный-файл.txt' in notes, 'list must show real files'
            else:
                assert state == 1 and plan, 'expected a reviewable plan'
                result = tool(':apply')
                assert result['state'] == 1, result['observation']
            actual = {p.relative_to(root).as_posix() for p in root.rglob('*') if p.is_file()}
            if label == 'excel-formula':
                import openpyxl
                books = list(root.rglob('*.xlsx'))
                assert len(books) == 1
                formulas = openpyxl.load_workbook(books[0], data_only=False)
                values = openpyxl.load_workbook(books[0], data_only=True)
                sums = [c.coordinate for row in formulas.active for c in row if c.data_type == 'f' and 'SUM(' in c.value]
                assert sums and any(values.active[address].value == 15 for address in sums), 'SUM formula or its correct Excel result missing'
            elif label in ('excel-path', 'excel-zero', 'excel-clarify', 'excel-word', 'word-there'):
                import openpyxl
                books, docs = list(root.rglob('*.xlsx')), list(root.rglob('*.docx'))
                assert len(books) == (0 if label == 'word-there' else 1)
                assert len(docs) == (1 if label in ('excel-word', 'word-there') else 0)
                for book in books:
                    wb = openpyxl.load_workbook(book)
                    assert wb.sheetnames == ['Лист1']
                    assert all(cell.value is None for row in wb.active for cell in row)
                    with zipfile.ZipFile(book) as archive:
                        assert archive.testzip() is None
                        for name in archive.namelist():
                            ET.fromstring(archive.read(name))
                if label in ('excel-path', 'excel-clarify'):
                    assert books[0].parent == root / 'Работа'
                elif label == 'excel-word':
                    assert books[0].parent == docs[0].parent == root / 'Офис'
                elif label == 'word-there':
                    assert docs[0].parent == root
                if label == 'excel-path':
                    # A second request uses the actual application state after apply,
                    # without repeating the folder or previous model conversation.
                    last = result['last']
                    assert Path(last) == books[0].parent
                    goal2 = 'создай там ворд'
                    tool(json.dumps({'goal': goal2}, ensure_ascii=False))
                    dialogue_notes = ''
                    for step in range(16):
                        prompt = (f'Задача пользователя: {goal2}\nПапки: @desktop={root}; '
                                  f'@downloads={root}\\Downloads; @documents={root}\\Documents; @last={last}.\n'
                                  + dialogue_notes)
                        body = json.dumps({'model': 'local', **schema, 'messages': [
                            {'role': 'system', 'content': system}, {'role': 'user', 'content': prompt}],
                            'temperature': 0.2, 'max_tokens': 512, 'stream': False}).encode('utf-8')
                        req = urllib.request.Request(f'http://127.0.0.1:{args.port}/v1/chat/completions', data=body,
                                                     headers={'Content-Type': 'application/json'})
                        with urllib.request.urlopen(req, timeout=180) as response:
                            answer = json.load(response)['choices'][0]['message']['content']
                        follow = tool(answer.replace('\n', ' '))
                        log.append({'followup': goal2, 'model': answer, 'result': follow})
                        dialogue_notes += f'\nВызов: {answer}\nРезультат: {follow["observation"]}\n'
                        if follow['state'] in (1, 2):
                            break
                    assert follow['state'] == 1 and len(follow['plan']) == 1
                    assert tool(':apply')['state'] == 1
                    docs2 = list(root.rglob('*.docx'))
                    assert len(docs2) == 1 and docs2[0].parent == books[0].parent
            elif label == 'word-stl':
                docs = list(root.rglob('*.docx'))
                assert len(docs) == 1, 'Word document omitted or duplicated'
                folder = docs[0].parent
                assert all((folder / name).is_file() and not (root / name).exists() for name in files if name.endswith('.stl'))
                assert (root / 'keep.stl.txt').is_file()
                with zipfile.ZipFile(docs[0]) as archive:
                    assert archive.testzip() is None
                    ET.fromstring(archive.read('[Content_Types].xml'))
                    ET.fromstring(archive.read('_rels/.rels'))
                    ET.fromstring(archive.read('word/document.xml'))
            elif label == 'inspect-34d':
                assert actual == set(files)
            else:
                assert actual == expected, f'actual={actual}; expected={expected}'
            for path in root.rglob('*'):
                if path.is_file() and path.suffix not in ('.docx', '.xlsx'):
                    assert path.read_text(encoding='utf-8') in files.values()
            report.append({'case': label, 'pass': True, 'seconds': round(time.time() - started, 2), 'log': log})
        except Exception as error:
            report.append({'case': label, 'pass': False, 'error': str(error), 'log': log})
        finally:
            process.stdin.close()
            process.wait(timeout=10)
    (args.scratch / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps([{k: v for k, v in row.items() if k != 'log'} for row in report], ensure_ascii=False, indent=2))
    if not all(r['pass'] for r in report):
        raise SystemExit(1)


if __name__ == '__main__':
    main()
