"""Real bundled QwenPaw + local GGUF + production native tool engine.

Fixtures and Office outputs stay strictly inside --scratch. No user files.
"""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import time
import zipfile


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    p = argparse.ArgumentParser()
    p.add_argument('--runtime', type=Path, required=True)
    p.add_argument('--probe', type=Path, required=True)
    p.add_argument('--scratch', type=Path, required=True)
    p.add_argument('--port', type=int, default=39285)
    p.add_argument('--case', action='append', help='Run selected cases after a targeted fix')
    p.add_argument('--keep-going', action='store_true', help='Report every selected case, including failures')
    a = p.parse_args()
    a.scratch.mkdir(parents=True, exist_ok=False)
    probe = str(a.probe.resolve())
    runtime = a.runtime.resolve()
    system = json.loads(subprocess.check_output([probe, '--system'], encoding='utf-8'))
    schema = json.loads(subprocess.check_output([probe, '--schema'], encoding='utf-8'))
    cases = [
        ('inventory', 'на рабочем столе создай папку и закинь туда ексели, и создай в нем ексель с перечислением того что в папке',
         {'a.xlsx': b'a', 'b.xls': b'b', 'c.xlsm': b'c', 'keep.xlsx.txt': b'keep', 'Модели/old.xlsx': b'old'}),
        ('word-stl', 'создай папку на рабочем столе и создай в нем ворд файл, и в ту же папку stl файлы с рабочего стола кинь',
         {'Деталь.stl': b'one', 'два.STL': b'two', 'keep.stl.txt': b'keep'}),
        ('read-34d', 'какие файлы лежат в папке 34д на рабочем столе?',
         {**{f'34д/Файл-{i:02}.txt': str(i).encode() for i in range(65)}, 'Модели/old.stl': b'wrong folder'}),
        ('excel-sum', 'Создай на рабочем столе Расчёт.xlsx и запиши числа 12 и 3, а рядом формулу SUM для их суммы.', {}),
        ('excel-table', 'Создай на рабочем столе Материалы.xlsx с оформленной таблицей: шапка Материал и Код; строки Сталь 40Х, 00123 и Алюминий, 00124. Сделай шапку жирной с цветом, добавь фильтр и закрепи первую строку. Коды оставь текстом с нулями.', {}),
        ('word-text', 'Создай на рабочем столе Отчёт.docx с текстом: Проверка локального помощника.', {}),
        ('mixed', 'В папке Результат на рабочем столе создай Word с текстом Привет и пустую книгу Excel, а все файлы .pdf с рабочего стола скопируй туда.',
         {'a.pdf': b'pdf-one', 'b.PDF': b'pdf-two', 'keep.pdf.txt': b'keep'}),
        ('unsupported', 'Удали все файлы на рабочем столе.', {'keep.txt': b'do not delete'}),
        ('dir-parent', 'создай в {desktop}\\Вайбкокинг папку и перенеси туда папку по пути "{desktop}\\вайбокок2"',
         {'Вайбкокинг/keep.txt': b'keep', 'вайбокок2/влож/данные.txt': b'nested data'}),
        ('dir-after', 'Создай папку в "{desktop}\\Вайбкокинг" и перенеси туда папку "{desktop}\\вайбокок2".',
         {'Вайбкокинг/keep.txt': b'keep', 'вайбокок2/влож/данные.txt': b'nested data'}),
        ('dir-direct', 'Перенеси папку "{desktop}\\вайбокок2" в "{desktop}\\Вайбкокинг".',
         {'Вайбкокинг/keep.txt': b'keep', 'вайбокок2/влож/данные.txt': b'nested data'}),
    ]
    report = []
    for label, goal, fixtures in cases:
        if a.case and label not in a.case: continue
        root = (a.scratch / label).resolve(); root.mkdir()
        goal = goal.replace('{desktop}', str(root))
        for name, data in fixtures.items():
            path = root / name; path.parent.mkdir(parents=True, exist_ok=True); path.write_bytes(data)
        if label.startswith('dir-'): (root/'вайбокок2/пустая').mkdir()
        process = subprocess.Popen([probe, str(root)], stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                   text=True, encoding='utf-8')
        def tool(action):
            process.stdin.write((action if isinstance(action, str) else json.dumps(action, ensure_ascii=False)) + '\n')
            process.stdin.flush()
            return json.loads(process.stdout.readline())
        result = tool({'goal': goal})
        notes, plan, log, new_dir = '', [], [], ''
        errors = 0
        state_file = root / 'backend-state.json'
        # Harness artifacts belong outside the simulated desktop.
        state_file = a.scratch.resolve() / (label + '.state.json')
        started = time.monotonic()
        try:
            for step in range(16):
                prompt = (f'Задача пользователя: {goal}\nПапки: @desktop={root}; @downloads={root / "Downloads"}; '
                          f'@documents={root / "Documents"}; @last={root}; @new={new_dir}.\n'
                          f'Шаг {step+1} из 16. Действий в плане: {len(plan)}.\n')
                kinds = ['mkdir', 'copy', 'move', 'rename', 'open', 'write', 'write', 'write', 'excel_write', 'inventory', 'move каталог']
                for action in plan[:8]:
                    prompt += f'В плане: {kinds[action["kind"]]} {action["src"]} -> {action["dst"]}\n'
                prompt += notes[-10000:]
                request = {'cursorpad': {'goal': result['model_goal'], 'plan': result['evidence'],
                    'observation': result['observation'], 'observation_state': result['state']}, 'model': 'local', **schema,
                           'messages': [{'role': 'system', 'content': system}, {'role': 'user', 'content': prompt}],
                           'temperature': 0.2, 'max_tokens': 512, 'stream': True}
                input_file = a.scratch.resolve() / 'request.json'
                output_file = a.scratch.resolve() / 'response.json'
                input_file.write_text(json.dumps(request, ensure_ascii=False), encoding='utf-8')
                subprocess.run([str(runtime/'python.exe'), '-I', '-B', str(runtime/'bridge.py'),
                    '--request', str(input_file), '--response', str(output_file), '--state', str(state_file),
                    '--endpoint', f'http://127.0.0.1:{a.port}/v1/chat/completions'], check=True, timeout=580)
                reply = json.loads(output_file.read_text(encoding='utf-8'))
                assert reply['ok'], reply
                action = json.loads(reply['action'])
                result = tool(action)
                plan, new_dir = result['plan'], result['new_dir']
                log.append({'action': action, 'result': result})
                print(label, step + 1, action, '->', result['state'], flush=True)
                notes += f'\nВызов: {reply["action"]}\nРезультат: {result["observation"]}\n'
                if result['state'] < 0:
                    errors += 1
                    assert errors < 3, 'production UI stops after 3 rejected tool calls'
                if result['state'] in (1, 2): break
            assert result['state'] in (1, 2), 'agent did not finish'
            assert {p.relative_to(root).as_posix(): p.read_bytes() for p in root.rglob('*') if p.is_file()} == fixtures
            state = json.loads(state_file.read_text(encoding='utf-8'))
            assert state['backend'] == 'QwenPaw' and state['requirements'] and state['agent_state']
            if label == 'unsupported':
                assert result['state'] == 2 and not plan
            elif label == 'read-34d':
                assert not plan and result['state'] == 1
                assert all(f'Файл-{i:02}.txt' in notes for i in range(65))
                assert 'old.stl' not in notes
            else:
                assert result['state'] == 1 and plan, 'expected a reviewable plan'
                applied = tool(':apply'); assert applied['state'] == 1, applied
                if label == 'inventory':
                    import openpyxl
                    path = Path([x for x in plan if x['kind'] == 9][0]['dst'])
                    rows = list(openpyxl.load_workbook(path).active.values)[2:]
                    assert {row[0] for row in rows} == {'a.xlsx', 'b.xls', 'c.xlsm'}
                    assert (root/'Модели/old.xlsx').read_bytes() == b'old'
                    for name in ('a.xlsx', 'b.xls', 'c.xlsm'):
                        assert not (root/name).exists() and (path.parent/name).read_bytes() == fixtures[name]
                elif label == 'word-stl':
                    assert len(list(root.rglob('*.docx'))) == 1
                    assert len([x for x in plan if x['kind'] == 2]) == 2
                    assert (root/'keep.stl.txt').read_bytes() == b'keep'
                elif label == 'excel-sum':
                    import openpyxl
                    book = openpyxl.load_workbook(next(root.glob('*.xlsx')), data_only=True)
                    assert any(cell.value == 15 for row in book.active for cell in row)
                elif label == 'word-text':
                    with zipfile.ZipFile(next(root.glob('*.docx'))) as z:
                        assert 'Проверка локального помощника.' in z.read('word/document.xml').decode()
                elif label == 'mixed':
                    dest = root/'Результат'
                    assert len(list(dest.glob('*.docx'))) == len(list(dest.glob('*.xlsx'))) == 1
                    for name in ('a.pdf', 'b.PDF'): assert (dest/name).read_bytes() == (root/name).read_bytes()
                elif label == 'excel-table':
                    import openpyxl
                    sheet = openpyxl.load_workbook(root/'Материалы.xlsx').active
                    assert sheet['A1'].font.bold and sheet.freeze_panes == 'A2' and sheet.auto_filter.ref == 'A1:B3'
                    assert list(sheet.values) == [('Материал', 'Код'), ('Сталь 40Х', '00123'), ('Алюминий', '00124')]
                elif label.startswith('dir-'):
                    assert len([x for x in plan if x['kind'] == 10]) == 1
                    dest = root/'Вайбкокинг'
                    if label != 'dir-direct': dest /= 'Новая папка'
                    dest /= 'вайбокок2'
                    assert (dest/'влож/данные.txt').read_bytes() == b'nested data'
                    assert (dest/'пустая').is_dir() and not (root/'вайбокок2').exists()
                    assert (root/'Вайбкокинг/keep.txt').read_bytes() == b'keep'
            report.append({'case': label, 'pass': True, 'seconds': round(time.monotonic()-started, 2),
                           'requirements': state['requirements'], 'skills': state['skills'], 'steps': log})
        except Exception as error:
            report.append({'case': label, 'pass': False, 'seconds':round(time.monotonic()-started, 2), 'error': str(error), 'steps': log})
            (a.scratch/'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
            if not a.keep_going: raise
        finally:
            process.stdin.close(); process.wait(timeout=10)
    (a.scratch/'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    passed = sum(row['pass'] for row in report)
    print('RESULT', passed, '/', len(report), 'real QwenPaw scenarios')
    if passed != len(report): raise SystemExit(1)


if __name__ == '__main__':
    main()
